"""Geometry and text-fit QA. Substitutes for render-based QA (no LibreOffice here).

Text-fit uses a conservative average-advance-width model: for Calibri/Cambria at size S pt,
mean glyph advance is ~0.50 em for mixed-case prose. Flags are advisory but catch the
overflow that render inspection is meant to catch.
"""
import sys, math
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.40          # hard minimum distance from slide edge
CHAR_EM = 0.50         # average advance width as fraction of font size
LINE_MULT = 1.22       # line box as fraction of font size
FILL = 1.04            # allow 4% slack before flagging

deck = Presentation(sys.argv[1])
issues = []

def shape_text_boxes(sl):
    out = []
    for sh in sl.shapes:
        if not sh.has_text_frame:
            continue
        txt = sh.text_frame.text
        if not txt.strip():
            continue
        sizes = [r.font.size.pt for p in sh.text_frame.paragraphs
                 for r in p.runs if r.font.size is not None]
        out.append((sh, txt, max(sizes) if sizes else 12.0))
    return out

for idx, sl in enumerate(deck.slides, 1):
    boxes = shape_text_boxes(sl)

    # 1. bounds
    for sh in sl.shapes:
        if sh.left is None or sh.top is None:
            continue
        x, y = sh.left/EMU, sh.top/EMU
        w = (sh.width or 0)/EMU
        h = (sh.height or 0)/EMU
        if x < -0.01 or y < -0.01 or x+w > SLIDE_W+0.01 or y+h > SLIDE_H+0.01:
            issues.append((idx, "OFFSLIDE",
                f"{sh.shape_type} at ({x:.2f},{y:.2f}) {w:.2f}x{h:.2f}"))

    # 2. text fit + margin
    for sh, txt, size in boxes:
        x, y = sh.left/EMU, sh.top/EMU
        w, h = sh.width/EMU, sh.height/EMU
        if x < MARGIN - 0.01 or x + w > SLIDE_W - MARGIN + 0.01:
            issues.append((idx, "MARGIN", f"text x-span {x:.2f}..{x+w:.2f}: {txt[:44]!r}"))
        if y < 0.30 or y + h > SLIDE_H - 0.22:
            issues.append((idx, "MARGIN", f"text y-span {y:.2f}..{y+h:.2f}: {txt[:44]!r}"))

        cpl = max(1, int(w * 72.0 / (CHAR_EM * size)))
        lines = 0
        for para in txt.split("\n"):
            lines += max(1, math.ceil(len(para) / cpl))
        need = lines * size * LINE_MULT / 72.0
        if need > h * FILL:
            issues.append((idx, "OVERFLOW",
                f"{size:.0f}pt in {w:.2f}x{h:.2f}\" needs ~{need:.2f}\" ({lines} lines): {txt[:50]!r}"))

    # 3. text-on-text overlap
    for i in range(len(boxes)):
        for j in range(i+1, len(boxes)):
            a, b = boxes[i][0], boxes[j][0]
            ax, ay, aw, ah = a.left/EMU, a.top/EMU, a.width/EMU, a.height/EMU
            bx, by, bw, bh = b.left/EMU, b.top/EMU, b.width/EMU, b.height/EMU
            ox = min(ax+aw, bx+bw) - max(ax, bx)
            oy = min(ay+ah, by+bh) - max(ay, by)
            if ox > 0.06 and oy > 0.06:
                issues.append((idx, "OVERLAP",
                    f"{ox:.2f}x{oy:.2f}\" — {boxes[i][1][:26]!r} / {boxes[j][1][:26]!r}"))

by_kind = {}
for slide, kind, msg in issues:
    by_kind.setdefault(kind, []).append((slide, msg))
print(f"{len(deck.slides)} slides checked; {len(issues)} issue(s)\n")
for kind in ("OFFSLIDE", "OVERFLOW", "OVERLAP", "MARGIN"):
    rows = by_kind.get(kind, [])
    print(f"--- {kind}: {len(rows)}")
    for slide, msg in rows[:40]:
        print(f"  slide {slide:>2}  {msg}")
